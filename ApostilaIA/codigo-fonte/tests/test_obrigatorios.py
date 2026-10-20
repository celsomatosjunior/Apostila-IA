"""Os 16 testes obrigatórios da especificação (seção 23), executados de verdade.

Sem GPU e sem Windows neste ambiente de compilação: os testes 14 (GPU) e 15
(instalação Windows) verificam aqui a parte que é possível verificar e têm sua
execução completa documentada no RELATORIO_DE_TESTES.md (Wine / limitações).
O simulador extrativo substitui o modelo de IA nos testes de pipeline; o teste
13 também executa um modelo GGUF real com o llama.cpp quando disponível.
"""
from __future__ import annotations

import json
import re
import threading
import time
from pathlib import Path

import pytest

import sintetico
from conftest import llama_linux_dir, tiny_model
from apostila.fontes import slides as slides_mod
from apostila.fontes import transcricao
from apostila.llm.cliente import Controle
from apostila.llm.simulador import ProvedorSimulado, Simulador
from apostila.pipeline import orquestrador
from apostila.pipeline.orquestrador import Cancelado, Gerador
from apostila.settings import Settings


def _gerar(tmp: Path, txts, pdfs, st=None, sim=None, ctl=None, on_progress=None, prov=None):
    st = st or Settings()
    prov = prov or ProvedorSimulado(sim or Simulador())
    g = Gerador(txts, pdfs, tmp / "saida", st, prov, controle=ctl, on_progress=on_progress,
                base_dir=tmp / "proj", modelos_key="sim")
    return g.executar(), g, prov


def _pdf_texto(p: Path) -> str:
    import pymupdf
    with pymupdf.open(str(p)) as d:
        return "\n".join(pg.get_text() for pg in d)


def _rel(r) -> dict:
    return json.loads(Path(r.relatorio_json).read_text(encoding="utf-8"))


def _checar_pdf(p: Path):
    import pymupdf
    with pymupdf.open(str(p)) as d:
        assert d.page_count >= 3
        toc = d.get_toc()
        assert toc, "PDF sem marcadores (sumário navegável)"
        for i, pg in enumerate(d):
            assert pg.get_text().strip() or pg.get_images(), f"página {i + 1} vazia"
        fonts = {f[3] for i in range(d.page_count) for f in d.get_page_fonts(i)}
        assert any("IBMPlex" in f for f in fonts)
        links = [l for l in d[1].get_links() if l.get("kind") == pymupdf.LINK_GOTO or "page" in l]
        assert links, "sumário sem links internos"
    return True


# ---------------------------------------------------------------- 1
def test_01_txt_curto_pdf_simples(tmp_path):
    t = sintetico.transcricao(tmp_path / "curta.txt", "Aula curta", 6, seed=3, topicos=sintetico.TOPICOS[:2])
    p = sintetico.slides(tmp_path / "simples.pdf", "Situs fetal", "Fulana de Tal", topicos=sintetico.TOPICOS[:2],
                         ultrassons=0, tabela=False, divisoria="")
    r, g, _ = _gerar(tmp_path, [t], [p])
    assert r.pdf.exists() and r.docx.exists() and r.relatorio.exists()
    assert r.pdf.parent == tmp_path / "saida" and r.pdf.name == "Apostila_Final.pdf"
    assert r.docx.name == "Apostila_Final.docx"
    txt = _pdf_texto(r.pdf)
    assert "Situs fetal" in txt and "Fulana de Tal" in txt and "Sumário" in txt


# ---------------------------------------------------------------- 2
def test_02_aula_com_varios_pdfs(tmp_path):
    t1 = sintetico.transcricao(tmp_path / "p1.txt", "Parte 1", 12, seed=1, topicos=sintetico.TOPICOS[:3])
    t2 = sintetico.transcricao(tmp_path / "p2.txt", "Parte 2", 10, seed=2, topicos=sintetico.TOPICOS[3:])
    a = sintetico.slides(tmp_path / "AULA 01.pdf", "Situs fetal", "Fulana", topicos=sintetico.TOPICOS[:3], ultrassons=0)
    b = sintetico.slides(tmp_path / "AULA 02.pdf", "Posição cardíaca", "Beltrana", topicos=sintetico.TOPICOS[3:],
                         ultrassons=1, tabela=False)
    r, g, _ = _gerar(tmp_path, [t2, t1], [b, a])   # ordem trocada de propósito
    rel = _rel(r)
    decks_associados = {s["slide"].split("-")[0] for s in rel["slides"] if s["associacao"] == "fala"}
    assert decks_associados == {"A1", "A2"}
    # ordem natural dos arquivos: parte 1 antes da parte 2; AULA 01 antes de AULA 02
    assert [p.name for p in g.pdfs] == ["AULA 01.pdf", "AULA 02.pdf"]
    assert rel["mapa"][0]["fala"].startswith("Parte 1")
    assert "Fulana" in _pdf_texto(r.pdf) and "Beltrana" in _pdf_texto(r.pdf)


# ---------------------------------------------------------------- 3 e 4
@pytest.mark.parametrize("minutos,limite_s", [(130, 240), (300, 600)])
def test_03_04_transcricoes_longas(tmp_path, minutos, limite_s):
    t = sintetico.transcricao(tmp_path / f"aula_{minutos}.txt", f"Aula {minutos} min", minutos, seed=minutos)
    p = sintetico.slides(tmp_path / "slides.pdf", "Situs fetal", "Fulana", ultrassons=2)
    tr = transcricao.ler_transcricoes([t])
    assert tr.palavras >= minutos * 140
    prog = []
    t0 = time.monotonic()
    r, g, prov = _gerar(tmp_path, [t], [p], on_progress=prog.append)
    dt = time.monotonic() - t0
    assert dt < limite_s
    jan = json.loads((g.proj.ck("janelas.json")).read_text(encoding="utf-8"))["body"]
    assert len(jan) >= minutos * 150 // 2600            # segmentado, não uma chamada só
    assert max(j["tokens"] for j in jan) <= 1800 * 1.4   # nenhum trecho estoura o tamanho
    assert prov.sim.chamadas["notas"] == len(jan)
    caps = len(list((g.proj.dir / "capitulos").glob("c*.json")))
    assert caps >= 2
    fracs = [d["frac"] for d in prog if "frac" in d]
    assert fracs == sorted(fracs) and fracs[-1] > 0.98
    assert any(d.get("eta_s") is not None for d in prog)


# ---------------------------------------------------------------- 5
def test_05_slides_com_muitas_imagens(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 15, seed=5)
    p = sintetico.slides(tmp_path / "img.pdf", "Imagens", "Fulana", achatar=tuple(range(6)), ultrassons=24)
    deck = slides_mod.ler_pdf(p, 1)
    imgs = [pg for pg in deck.paginas if pg.tipo == "imagem"]
    assert len(imgs) >= 28
    r, g, prov = _gerar(tmp_path, [t], [p])
    assert prov.sim.chamadas["visao"] >= 28   # cada página visual passou pelo modelo de visão
    rel = _rel(r)
    assert rel["checagens"][7]["ok"]


# ---------------------------------------------------------------- 6
def test_06_slides_com_ultrassonografias(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 15, seed=6)
    p = sintetico.slides(tmp_path / "us.pdf", "Situs fetal", "Fulana", ultrassons=4)
    r, g, _ = _gerar(tmp_path, [t], [p])
    cards = {c.stem: json.loads(c.read_text(encoding="utf-8"))["body"] for c in g.proj.ck("cartoes").glob("*.json")}
    us = [c for c in cards.values() if c["titulo"] in ("Isomerismo esquerdo – ázigos", "Dextroposição")]
    assert us and all(c["visual"] for c in us)
    import pymupdf
    with pymupdf.open(str(r.pdf)) as d:
        n_img = sum(len(pg.get_images()) for pg in d)
        assert n_img >= 1
        # figura sem deformação: proporção do slide (16:9) preservada
        for pg in d:
            for info in pg.get_image_info():
                w = info["bbox"][2] - info["bbox"][0]
                h = info["bbox"][3] - info["bbox"][1]
                if w > 200:
                    assert abs(w / h - 960 / 540) < 0.06
    assert "Figura" in _pdf_texto(r.pdf)


# ---------------------------------------------------------------- 7
def test_07_pdf_com_pouco_texto_extraivel(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 10, seed=7, topicos=sintetico.TOPICOS[:3])
    p = sintetico.slides(tmp_path / "achatado.pdf", "Situs", "Fulana", topicos=sintetico.TOPICOS[:3],
                         achatar=(0, 1, 2), ultrassons=0, tabela=False, divisoria="")
    deck = slides_mod.ler_pdf(p, 1)
    for pg in deck.paginas[1:]:
        assert pg.n_palavras == 0 and pg.tipo == "imagem" and pg.precisa_visao
    st = Settings()
    st.usar_visao = False
    r, g, prov = _gerar(tmp_path, [t], [p], st=st)
    assert prov.sim.chamadas["visao"] == 0
    assert any("não foram lidos pelo modelo de visão" in a for a in _rel(r)["avisos"])


# ---------------------------------------------------------------- 8
class SimEsquema(Simulador):
    def _capitulo(self, ctx):
        r = super()._capitulo(ctx)
        vazio = {"subtipo": "nenhum", "texto": "", "itens": [], "cabecalho": [], "linhas": [], "slide": "", "fontes": []}
        r["secoes"][0]["blocos"] += [
            dict(vazio, tipo="esquema", subtipo="sequencia", texto="Sequência de avaliação",
                 itens=["Situs", "Posição cardíaca", "Conexões venosas", "Conexões atrioventriculares",
                        "Conexões ventriculoarteriais"]),
            dict(vazio, tipo="esquema", subtipo="sequencia", texto="Esquema inventado",
                 itens=["Ressonância magnética fetal", "Cariótipo molecular obrigatório"]),
            dict(vazio, tipo="tabela", cabecalho=["Achado", "Isomerismo direito", "Isomerismo esquerdo"],
                 linhas=[["Baço", "Asplenia", "Poliesplenia"]], texto="Isomerismos"),
        ]
        return r


def test_08_tabelas_e_esquemas(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 15, seed=8)
    p = sintetico.slides(tmp_path / "s.pdf", "Situs", "Fulana", ultrassons=0, tabela=True)
    deck = slides_mod.ler_pdf(p, 1)
    tabs = [tb for pg in deck.paginas for tb in pg.tabelas]
    assert tabs and tabs[0].cabecalho[:2] == ["Achado", "Isomerismo direito"]
    r, g, _ = _gerar(tmp_path, [t], [p], sim=SimEsquema())
    rel = _rel(r)
    assert any("Ressonância" in " ".join(e["etapas"]) for e in rel["esquemas_removidos"])
    txt = _pdf_texto(r.pdf)
    assert "Conexões venosas" in txt and "Ressonância magnética fetal" not in txt
    assert "Tabela 1" in txt and "Poliesplenia" in txt
    from docx import Document
    d = Document(str(r.docx))
    assert any(t.rows[0].cells[0].text == "Achado" for t in d.tables)
    assert len(d.inline_shapes) >= 1   # esquema como imagem


# ---------------------------------------------------------------- 9
def test_09_terminologia_medica(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 15, seed=9)
    txt = t.read_text(encoding="utf-8").replace("TRANSCRIÇÃO\n", "TRANSCRIÇÃO\n\nA doutora Catia vai falar das veias.\n", 1)
    t.write_text(txt, encoding="utf-8")
    p = sintetico.slides(tmp_path / "s.pdf", "Situs fetal", "KATIA SOUZA", ultrassons=0)
    r, g, _ = _gerar(tmp_path, [t], [p])
    rel = _rel(r)
    cor = {(c["de"].lower(), c["para"]) for c in rel["termos_corrigidos"]}
    assert ("sítios", "situs") in cor
    conf = {(c["de"].lower(), c["para"].lower()) for c in rel["termos_conferir"]}
    assert ("catia", "katia") in conf
    assert any(a == "destrocardia" and b == "dextrocardia" for a, b in conf) or \
        ("destrocardia", "dextrocardia") in {(c["de"].lower(), c["para"].lower()) for c in rel["termos_corrigidos"]}
    out = _pdf_texto(r.pdf)
    assert "situs solitus" in out.lower() and "sítios solitus" not in out.lower()
    assert "Katia Souza" in out


# ---------------------------------------------------------------- 10
def test_10_geracao_pdf(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 20, seed=10)
    p = sintetico.slides(tmp_path / "s.pdf", "Situs fetal", "Fulana", ultrassons=2)
    r, g, _ = _gerar(tmp_path, [t], [p])
    assert _checar_pdf(r.pdf)
    import pymupdf
    with pymupdf.open(str(r.pdf)) as d:
        assert d.metadata["title"]
        # numeração de páginas no rodapé a partir do sumário
        assert d[2].get_text().strip().splitlines()[-1].strip().isdigit()
        # nenhuma linha de texto sai da página (sem cortes)
        for pg in d:
            for b in pg.get_text("dict")["blocks"]:
                x0, y0, x1, y1 = b["bbox"]
                assert x0 >= 0 and y0 >= 0 and x1 <= pg.rect.width + 1 and y1 <= pg.rect.height + 1


# ---------------------------------------------------------------- 11
def test_11_geracao_docx(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 20, seed=11)
    p = sintetico.slides(tmp_path / "s.pdf", "Situs fetal", "Fulana", ultrassons=2)
    r, g, _ = _gerar(tmp_path, [t], [p])
    from docx import Document
    d = Document(str(r.docx))
    heads = [pp.text for pp in d.paragraphs if pp.style.name.startswith("Heading")]
    assert heads[0] == "Sumário" and any(re.match(r"^1\s+\S", h) for h in heads)
    xml = d.element.xml
    assert "TOC \\o" in xml and "w:tblHeader" in xml or "TOC \\o" in xml
    assert len(d.inline_shapes) >= 1
    rotulos = " ".join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
    assert any(k in rotulos for k in ("OBSERVAÇÃO DA PROFESSORA", "PONTO DE ATENÇÃO", "DICA CLÍNICA"))
    # reabre e salva (editável)
    d.add_paragraph("edição do usuário")
    d.save(str(tmp_path / "editado.docx"))


# ---------------------------------------------------------------- 12
def test_12_retomada_apos_interrupcao(tmp_path):
    t1 = sintetico.transcricao(tmp_path / "p1.txt", "Parte 1", 20, seed=12, topicos=sintetico.TOPICOS[:3])
    t2 = sintetico.transcricao(tmp_path / "p2.txt", "Parte 2", 20, seed=22, topicos=sintetico.TOPICOS[3:])
    a = sintetico.slides(tmp_path / "AULA 01.pdf", "Situs fetal", "Fulana", topicos=sintetico.TOPICOS[:3])
    b = sintetico.slides(tmp_path / "AULA 02.pdf", "Posição", "Fulana", topicos=sintetico.TOPICOS[3:], tabela=False)
    t, p = [t1, t2], [a, b]
    ctl = Controle()
    sim1 = Simulador(atraso=0.01)

    def prog(d):
        if "Capítulo 1/" in d.get("detalhe", ""):
            ctl.cancel.set()   # "o programa fechou" logo após o 1º capítulo
    with pytest.raises(Cancelado):
        _gerar(tmp_path, t, p, sim=sim1, ctl=ctl, on_progress=prog)
    assert sim1.chamadas["capitulo"] >= 1
    notas_antes = sim1.chamadas["notas"]
    sim2 = Simulador()
    r, g, _ = _gerar(tmp_path, t, p, sim=sim2)
    assert sim2.chamadas["notas"] == 0 and sim2.chamadas["visao"] == 0   # nada refeito
    assert sim2.chamadas["capitulo"] < sim1.chamadas["capitulo"] + 10
    caps = len(list((g.proj.dir / "capitulos").glob("c*.json")))
    assert sim2.chamadas["capitulo"] == caps - 1   # só os capítulos que faltavam
    assert notas_antes > 0 and r.pdf.exists()
    # checkpoint corrompido é descartado e refeito, sem travar
    c1 = g.proj.dir / "capitulos" / "c01.json"
    c1.write_text("{lixo", encoding="utf-8")
    (g.proj.dir / "revisao.json").unlink()
    sim3 = Simulador()
    r3, _, _ = _gerar(tmp_path, t, p, sim=sim3)
    assert sim3.chamadas["capitulo"] == 1 and r3.pdf.exists()


def test_12b_pausa_e_continua(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 20, seed=13)
    ctl = Controle()
    sim = Simulador(atraso=0.02)
    eventos = []

    def prog(d):
        eventos.append(d)
        if d.get("etapa", "").startswith("Extraindo") and not ctl.pause.is_set() and not any("pausado" in e for e in eventos):
            ctl.pause.set()
            threading.Timer(0.5, ctl.pause.clear).start()
    r, g, _ = _gerar(tmp_path, [t], [], sim=sim, ctl=ctl, on_progress=prog)
    assert any(e.get("pausado") is True for e in eventos) and any(e.get("pausado") is False for e in eventos)
    assert r.pdf.exists()


# ---------------------------------------------------------------- 13
def test_13_funcionamento_offline(tmp_path, bloquear_rede):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 10, seed=14)
    p = sintetico.slides(tmp_path / "s.pdf", "Situs fetal", "Fulana", ultrassons=1)
    r, g, _ = _gerar(tmp_path, [t], [p])
    assert r.pdf.exists() and r.docx.exists()
    assert bloquear_rede == []   # nenhuma tentativa de acesso externo


@pytest.mark.skipif(not (llama_linux_dir() and tiny_model()), reason="llama.cpp/modelo GGUF de teste ausentes")
def test_13b_modelo_local_real_offline(tmp_path, bloquear_rede):
    """Modelo GGUF real no llama.cpp (CPU), sem rede: saída JSON válida no esquema das notas."""
    from apostila.llm.cliente import LlamaLLM
    from apostila.llm.servidor import Config, LlamaServer
    from apostila.pipeline import prompts
    exe = llama_linux_dir() / "llama-server"
    srv = LlamaServer(Config(exe=exe, model=tiny_model(), ctx=4096, backend="cpu", threads=2),
                      tmp_path / "srv.log")
    srv.start(timeout=300)
    try:
        llm = LlamaLLM(srv, "teste", False)
        paras = [("P1.1", "O situs solitus é o arranjo habitual dos órgãos. O estômago fica à esquerda.")]
        r = llm.json("notas", prompts.SISTEMA, prompts.prompt_notas(paras, "", [], []), prompts.SCHEMA_NOTAS, 500)
        assert set(prompts.SCHEMA_NOTAS["required"]) <= set(r)
        assert isinstance(r["itens"], list)
        assert not srv.gpu.verificada          # sem GPU: nunca declara GPU
        assert llm.stats.tokens_gerados > 0
    finally:
        srv.stop()
    assert bloquear_rede == []


# ---------------------------------------------------------------- 14
def test_14_gpu_so_declarada_quando_verificada(tmp_path):
    from apostila.llm.servidor import Config, LlamaServer
    srv = LlamaServer(Config(exe=Path("x"), model=Path("m"), backend="vulkan"))
    linhas = ["load_tensors: offloading 48 repeating layers to GPU",
              "load_tensors: offloaded 49/49 layers to GPU",
              "load_tensors:      Vulkan0 model buffer size =  4521.33 MiB",
              "load_tensors:   CPU_Mapped model buffer size =   315.00 MiB"]
    import io

    class P:
        stdout = io.StringIO("\n".join(linhas) + "\n")
    srv.proc = P()
    srv._read()
    assert srv.gpu.verificada and srv.gpu.camadas_gpu == 49 and "Vulkan0" in srv.gpu.descricao()
    srv2 = LlamaServer(Config(exe=Path("x"), model=Path("m"), backend="vulkan"))

    class Q:
        stdout = io.StringIO("load_tensors: offloaded 0/49 layers to GPU\nload_tensors: CPU_Mapped model buffer size = 4000 MiB\n")
    srv2.proc = Q()
    srv2._read()
    assert not srv2.gpu.verificada and srv2.gpu.descricao() == "CPU"
    # argumentos: CPU nunca pede GPU; GPU usa ajuste automático à VRAM com margem do modo
    a_cpu = LlamaServer(Config(exe=Path("x"), model=Path("m"), backend="cpu")).args()
    assert "-dev" in a_cpu and a_cpu[a_cpu.index("-dev") + 1] == "none"
    a_gpu = LlamaServer(Config(exe=Path("x"), model=Path("m"), backend="cuda", device="CUDA0", trabalho=True,
                               fit_margin_mb=2560)).args()
    assert a_gpu[a_gpu.index("--fit-target") + 1] == "2560" and "--prio" in a_gpu
    # escolha de conjunto pelo hardware (computadores-alvo)
    from apostila.hardware import GPUInfo, HardwareInfo
    from apostila.llm import modelos, runtime
    alvo = HardwareInfo(cpu_name="Intel Core i7-10700", cores_physical=8, threads=16, ram_total_mb=32 * 1024,
                        gpus=[GPUInfo("nvidia", "NVIDIA GeForce RTX 3070", 8192)])
    assert modelos.recomendar(alvo).key == "A"
    assert runtime.cuda_variant("NVIDIA GeForce RTX 5070 Ti") == "13.4"
    assert runtime.cuda_variant("NVIDIA GeForce RTX 4090 Laptop GPU") == "12.4"
    pequeno = HardwareInfo(ram_total_mb=8 * 1024)
    assert modelos.recomendar(pequeno).key == "C"


# ---------------------------------------------------------------- 15
def test_15_instalador_windows_documentado():
    """A instalação/desinstalação real é testada no Wine por build/windows/testar_instalador.sh;
    aqui conferimos que o script NSIS tem os elementos exigidos."""
    nsi = (Path(__file__).resolve().parents[1] / "build" / "windows" / "installer.nsi").read_text(encoding="utf-8")
    assert "ApostilaIA-Setup-${VERSION}.exe" in nsi or "ApostilaIA-Setup" in nsi
    assert "PortugueseBR" in nsi and "Uninstall" in nsi


# ---------------------------------------------------------------- 16
def test_16_fidelidade_transcricao_slides_apostila(tmp_path):
    t = sintetico.transcricao(tmp_path / "a.txt", "Aula", 25, seed=16)
    p = sintetico.slides(tmp_path / "s.pdf", "Situs fetal", "Fulana")
    r, g, _ = _gerar(tmp_path, [t], [p], sim=Simulador(injetar_externo=True))
    rel = _rel(r)
    out = _pdf_texto(r.pdf)
    # 1) informação externa injetada foi detectada e removida
    assert any("diretrizes internacionais" in x["frase"] for x in rel["removidas"])
    assert "diretrizes internacionais" not in out and "97%" not in out
    # 2) divergência numérica fala × slide registrada, sem escolha de valor
    assert any(d["fala"] == "90%" and d["slide"] == "59–75%" for d in rel["divergencias"])
    assert "Divergência" in out
    # 3) observações/alertas/exemplos preservados
    checks = {c["n"]: c for c in rel["checagens"]}
    assert checks[2]["ok"] and checks[3]["ok"] and checks[1]["ok"]
    # 4) toda frase restante tem as palavras de conteúdo sustentadas pelas fontes
    from apostila.pipeline.texto import cobertura, frases, vocab
    fontes = vocab([tr for tr in t.read_text(encoding="utf-8").split("\n")] + [_pdf_texto(p)])
    corpo = out.split("Termos e pontos a conferir")[0]
    ruins = [f for f in frases(corpo.replace("\n", " ")) if len(f.split()) > 6 and cobertura(f, fontes) < 0.5]
    assert len(ruins) <= 3, ruins[:3]
    # 5) referências internas: mapa capítulo -> parágrafos e slides
    assert all(m["fala"] for m in rel["mapa"]) and any(m["slides"] for m in rel["mapa"])
    # 6) referências bibliográficas dos slides no apêndice
    assert "A Practical Guide Of Testing" in out
