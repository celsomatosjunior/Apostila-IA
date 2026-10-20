"""DOCX editável no Microsoft Word (python-docx).

- Títulos com os estilos nativos "Título 1/2" (painel de navegação do Word e sumário).
- Sumário como campo TOC: o Word atualiza ao abrir (updateFields).
- Quadros de destaque como tabela de 1 célula sombreada (editável, não imagem).
- Tabelas com linha de cabeçalho repetida em cada página.
- Figuras dos slides em alta qualidade; esquemas como imagem nítida (220 ppi).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .. import logs
from .comum import preparar

log = logs.get("docx")

AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL2 = RGBColor(0x2E, 0x75, 0xB6)
CINZA = RGBColor(0x5B, 0x64, 0x70)
CORPO = "Cambria"
TITULOS = "Calibri"
QUADROS = {
    "observacao": ("EDF4FB", "2E75B6", "OBSERVAÇÃO DA PROFESSORA"),
    "dica": ("EDF7F0", "2E7D4F", "DICA CLÍNICA"),
    "atencao": ("FDF4E4", "B86E00", "PONTO DE ATENÇÃO"),
    "exemplo": ("F2F0F8", "5E4FA2", "EXEMPLO DISCUTIDO"),
    "passo_a_passo": ("EAF5F5", "1C7C7D", "PASSO A PASSO"),
}


def _rich(par, texto: str, size=None, color=None, font=None, bold=False):
    for i, parte in enumerate(str(texto).split("**")):
        if not parte:
            continue
        r = par.add_run(parte)
        r.bold = bold or (i % 2 == 1)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if font:
            r.font.name = font
    return par


def _shade(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _borders(cell, left_color: str, left_sz: int = 24, others: str = "nil"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), others)
        b.append(e)
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(left_sz))
    e.set(qn("w:color"), left_color)
    b.append(e)
    tcPr.append(b)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.append(c)


def _field(par, instr: str, placeholder: str = ""):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    r._r.append(f1)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    r3._r.append(f2)
    if placeholder:
        par.add_run(placeholder)
    r4 = par.add_run()
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    r4._r.append(f3)


def _update_fields_on_open(doc):
    settings = doc.settings.element
    u = OxmlElement("w:updateFields")
    u.set(qn("w:val"), "true")
    settings.append(u)


def _estilos(doc):
    st = doc.styles
    n = st["Normal"]
    n.font.name = CORPO
    n.font.size = Pt(11)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), CORPO)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 18), (2, 13.5), (3, 12)):
        h = st[f"Heading {lvl}"]
        h.font.name = TITULOS
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = AZUL
        rfonts = h.element.rPr.rFonts
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(att)) is not None:
                del rfonts.attrib[qn(att)]
        rfonts.set(qn("w:ascii"), TITULOS)
        rfonts.set(qn("w:hAnsi"), TITULOS)
        h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    cap = st["Caption"]
    cap.font.name = TITULOS
    cap.font.size = Pt(9)
    cap.font.italic = False
    cap.font.color.rgb = CINZA


def _quadro(doc, b):
    fill, cor, rot = QUADROS.get(b.get("subtipo"), QUADROS["observacao"])
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    _shade(c, fill)
    _borders(c, cor)
    p = c.paragraphs[0]
    r = p.add_run(rot)
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.name = TITULOS
    r.font.color.rgb = RGBColor.from_string(cor)
    if b.get("texto"):
        _rich(c.add_paragraph(), b["texto"], size=10.5)
    for k, it in enumerate(b.get("itens", []), 1):
        pp = c.add_paragraph()
        pp.paragraph_format.left_indent = Cm(0.5)
        pp.paragraph_format.space_after = Pt(2)
        pref = f"{k}. " if b.get("subtipo") == "passo_a_passo" else "• "
        pp.add_run(pref)
        _rich(pp, it, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _tabela(doc, b, ntab):
    if b.get("texto"):
        cap = doc.add_paragraph(style="Caption")
        cap.add_run(f"Tabela {ntab} — {b['texto']}")
        cap.paragraph_format.keep_with_next = True
    ncol = len(b["cabecalho"])
    t = doc.add_table(rows=1 + len(b["linhas"]), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(b["cabecalho"]):
        c = t.cell(0, j)
        _shade(c, "1F4E79")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = TITULOS
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _repeat_header(t.rows[0])
    for i, linha in enumerate(b["linhas"], 1):
        _cant_split(t.rows[i])
        for j in range(ncol):
            c = t.cell(i, j)
            if i % 2 == 0:
                _shade(c, "F4F7FA")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _rich(p, linha[j] if j < len(linha) else "", size=9.5, font=TITULOS)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def gerar(apost: dict, saida: Path, work: Path) -> Path:
    from .typst_pdf import esquema_png

    d = preparar(apost, work)
    doc = Document()
    _estilos(doc)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
    sec.different_first_page_header_footer = True

    # ---- capa
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Cm(3)
    _rich(p, (d.get("curso") or "Apostila de estudo").upper(), size=10, color=AZUL2, font=TITULOS, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Cm(1.5)
    _rich(p, d["titulo"], size=28, color=AZUL, font=TITULOS, bold=True)
    if d["professores"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Cm(0.8)
        _rich(p, ("Professora: " if len(d["professores"]) == 1 else "Professoras: ") + ", ".join(d["professores"]),
              size=12, color=CINZA, font=TITULOS)
    if len(d["aulas"]) > 1:
        for a in d["aulas"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            _rich(p, f"• {a['rotulo']}" + (f" — {a['professor']}" if a.get("professor") else ""), size=10, color=CINZA, font=TITULOS)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Cm(8)
    _rich(p, "Apostila de estudo elaborada exclusivamente a partir da transcrição da aula e dos slides apresentados. "
             f"Perfil: {d['perfil_rotulo']}. Gerada em {d['data']} pelo ApostilaIA {d['versao']} (processamento local).",
          size=9, color=CINZA, font=TITULOS)
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---- cabeçalho e rodapé (a partir da 2ª página)
    hp = sec.header.paragraphs[0]
    _rich(hp, d["titulo_curto"], size=8, color=CINZA, font=TITULOS)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE", "1")

    # ---- sumário
    h = doc.add_paragraph("Sumário", style="Heading 1")
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-2" \\h \\z \\u', "Clique com o botão direito e escolha “Atualizar campo” para montar o sumário.")
    p.add_run().add_break(WD_BREAK.PAGE)
    _update_fields_on_open(doc)

    nfig = ntab = 0
    for cap in d["capitulos"]:
        hh = doc.add_paragraph(f"{cap['num']}  {cap['titulo']}", style="Heading 1")
        if cap["num"] > 1:
            hh.paragraph_format.page_break_before = True
        ns = 0
        for s in cap["secoes"]:
            if s["titulo"]:
                ns += 1
                doc.add_paragraph(f"{cap['num']}.{ns}  {s['titulo']}", style="Heading 2")
            for b in s["blocos"]:
                t = b["tipo"]
                if t == "paragrafo":
                    _rich(doc.add_paragraph(), b["texto"])
                elif t == "topicos":
                    if b.get("texto"):
                        _rich(doc.add_paragraph(), b["texto"])
                    for it in b["itens"]:
                        _rich(doc.add_paragraph(style="List Bullet"), it)
                elif t == "quadro":
                    _quadro(doc, b)
                elif t == "tabela":
                    ntab += 1
                    _tabela(doc, b, ntab)
                elif t == "figura":
                    nfig += 1
                    pp = doc.add_paragraph()
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pp.paragraph_format.keep_with_next = True
                    pp.add_run().add_picture(b["arquivo_abs"], width=Cm(15.5))
                    cp = doc.add_paragraph(style="Caption")
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.add_run(f"Figura {nfig} — {b['texto']} ").font.size = Pt(9)
                    r = cp.add_run(f"({b['origem']})")
                    r.italic = True
                    r.font.size = Pt(9)
                elif t == "esquema":
                    png = work / f"esquema_{cap['num']}_{ns}_{nfig}_{ntab}_{id(b)}.png"
                    try:
                        esquema_png(b, work / "esq", png)
                        pp = doc.add_paragraph()
                        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        pp.add_run().add_picture(str(png), width=Cm(15.5))
                    except Exception as e:  # noqa: BLE001 — sem a imagem, o esquema vira lista editável
                        log.warning("Esquema como imagem falhou (%s); usando lista", e)
                        if b.get("texto"):
                            _rich(doc.add_paragraph(), f"**{b['texto']}**")
                        sep = " → " if b.get("subtipo") == "sequencia" else "\n"
                        _rich(doc.add_paragraph(), sep.join(b["itens"]))
            if d.get("mostrar_fontes") and s.get("fontes_rotulo"):
                pp = doc.add_paragraph()
                _rich(pp, "Fontes: " + s["fontes_rotulo"], size=7.5, color=CINZA, font=TITULOS)

    if d.get("pontos_conferir"):
        h = doc.add_paragraph("Termos e pontos a conferir", style="Heading 1")
        h.paragraph_format.page_break_before = True
        _rich(doc.add_paragraph(), "Itens que a revisão automática não pôde confirmar pelas fontes. Não foram corrigidos por suposição.",
              size=9.5, color=CINZA)
        for pt in d["pontos_conferir"]:
            pp = doc.add_paragraph(style="List Bullet")
            if pt["tipo"] == "divergencia":
                _rich(pp, "**Divergência:** ", size=10)
            _rich(pp, pt["texto"], size=10)
    if d.get("referencias"):
        h = doc.add_paragraph("Referências citadas nos slides", style="Heading 1")
        h.paragraph_format.page_break_before = not d.get("pontos_conferir")
        _rich(doc.add_paragraph(), "Obras e artigos impressos nos slides da aula (o conteúdo da apostila não foi complementado com elas).",
              size=9.5, color=CINZA)
        for r in d["referencias"]:
            _rich(doc.add_paragraph(style="List Bullet"), r, size=10)

    cp = doc.core_properties
    cp.title = d["titulo"]
    cp.author = ", ".join(d["professores"]) or "ApostilaIA"
    cp.comments = f"Gerado localmente pelo ApostilaIA {d['versao']}"
    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(saida))
    log.info("DOCX gerado: %s", saida)
    return saida
